# -*- coding: utf-8 -*-
"""부서 이음표 — 사역 임명 통합 엑셀(2010~2026) → 확인용 엑셀 (2026-10-06 · 이 PC 에서만)

  python tools/history/dept_lineage_draft.py [--out 결과.xlsx] [명단 엑셀 ...]
    기본 명단: C:/Projects/Data/정리/2010-2021_사역임명_통합.xlsx · C:/Projects/Data/2022-2026_사역임명_통합.xlsx
    기본 결과: C:/Projects/Data/정리/부서이음표_2차.xlsx

무엇: 해마다 달리 적힌 부서·팀 이름(부서 97가지 · 부서-팀 쌍 541가지)을 **큰 분류 → 계열 → 표준 팀** 에 잇는 표.
      사역 통계(시계열·부서)를 내기 전에 사역 담당과 친구가 이 엑셀을 보고 확인한다.
      설계·근거: 성경암송 저장소 docs/analysis/2026-10-06-ministry-stats-options.md · 2026-10-06-work-ministry-dept-lineage-draft.md
친구 결정(2026-10-06): 큰 분류는 찬양 · 교회학교 · 그 밖 — 목양(L-12 등)·기관(여전도회 등) 줄은 따로 · 은빛시니어학교·장학은 그 밖 ·
      부서는 해마다 달라지고 통폐합이 있으니 잇는다 · 교구는 보지 않는다.
1차 → 2차(2026-10-06): 1차 엑셀의 「확인해 주실 것」 17가지에 친구가 답했다(ANSWER — 답 적힌 파일은 부서이음표_초안_답_2026-10-06.xlsx).
      크게 바뀐 것 넷 — ① 「부설기관」 계열을 만들고 복지재단·봉사센터를 거기로 합쳤다(도서관 · 시니어학교 · 샬롬부·사랑부 포함)
      ② 샬롬부·사랑부를 교회학교에서 뺐다 ③ 찬양부 아래 중보기도팀을 전도·중보기도로 ④ 부서 칸이 빈 「주보간지」를 예배로.

⚠️ 명단도 결과도 저장소 밖에 둔다. 이 파일에는 부서·팀 이름과 규칙만 있다(사람 이름 없음).
   사람 이름은 「앞 해 그 팀에 있던 분이 뒤 해 그 팀에 몇 분 있나」를 세는 데만 쓰고, 결과에는 **수만** 적는다.
⚠️ 사람 수는 어림이다 — 한 해 안에서는 (이름, 목장)이 같으면 한 분, 앞뒤 해는 이름이 같으면 같은 분으로 센다(동명이인이 섞인다).
   운영에서는 교인ID 로 다시 센다.
⚠️ 규칙을 고칠 때: 부서를 이름만 보고 잇지 말 것 — 아래 EVENTS 처럼 **팀까지 내려가** 앞뒤 해 사람 흐름으로 확인한다.
⚠️ 친구가 답을 적어 돌려준 엑셀 위에 다시 만들지 말 것(답이 지워진다) — 받은 파일은 이름을 바꿔 두고, 새 판은 다른 이름으로 낸다.
"""
import collections, os, re, sys, unicodedata
import openpyxl
import openpyxl.worksheet.properties
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DEFAULT_IN = [r"C:\Projects\Data\정리\2010-2021_사역임명_통합.xlsx", r"C:\Projects\Data\2022-2026_사역임명_통합.xlsx"]
DEFAULT_OUT = r"C:\Projects\Data\정리\부서이음표_2차.xlsx"
ANSWERED_ON = "2026-10-06"

G_PRAISE, G_SCHOOL, G_ETC, G_MOK, G_ORG = "찬양", "교회학교", "그 밖", "따로 · 목양", "따로 · 기관"
GROUPS = [G_PRAISE, G_SCHOOL, G_ETC, G_MOK, G_ORG]


def nfc(v):
    return unicodedata.normalize("NFC", "" if v is None else str(v)).strip()


def ns(s):
    return re.sub(r"\s+", "", s)


# ── 부서 이름 바로잡기(띄어쓰기·쉼표·오타) ───────────────────────────────
DEPT_FIX = {
    "섬교부": "선교부", "세물결부": "새물결부", "소포츠운영부": "스포츠운영부", "차": "차량부", "예배": "예배부",
    "희망복지재단": "희망의복지재단", "멀티미디어": "멀티미디어운영부", "멀티미디어운영": "멀티미디어운영부",
    "감사위원": "감사위원회",
}


def dept_key(d):
    k = ns(d).replace(",", "")
    return DEPT_FIX.get(k, k)


# ── 팀 이름 바로잡기 ─────────────────────────────────────────────────
TEAM_TYPO = {
    "1-2학견부": "1-2학년부", "1-2학년": "1-2학년부", "5-6학녀부": "5-6학년부", "꿈너무학교": "꿈너머학교", "까페운영": "카페운영",
    "자원봉자지원": "자원봉사지원", "희망푸트뱅크": "희망푸드뱅크", "금요성령집화찬양": "금요성령집회찬양", "주알찬양3부": "주일찬양3부",
    "정책관리": "정착관리",
}
PREFIX_RE = re.compile(r"^(미취학|아동|청소년|사역팀|장년|사역지원)-")
PREFIX_NOHYPHEN_RE = re.compile(r"^(미취학|아동|청소년)(?=(사랑부|중등부|고등부|유년|소년|영아|유아|유치|새싹))")


def team_key(t):
    k = TEAM_TYPO.get(ns(t), ns(t))
    k = PREFIX_RE.sub("", k)
    k = PREFIX_NOHYPHEN_RE.sub("", k)
    return TEAM_TYPO.get(k, k)


# ── 부서 → 계열(그 밖) 기본값. 팀 이름이 정하는 것은 classify 가 먼저 가른다 ──────────────
F_EVAN, F_SERVE, F_AFFIL, F_WORSHIP, F_FACIL, F_ADMIN = "전도·중보기도", "교회봉사", "부설기관", "예배", "시설·안전", "총무·홍보"
F_NEW, F_DISC, F_MISSION, F_CAR, F_MEDIA, F_CULT = "새가족", "제자양육·가정사역", "선교", "차량", "방송·전산", "문화·스포츠"
F_FIN, F_AUDIT, F_SCHOLAR, F_ENDED, F_UNKNOWN = "재정", "감사", "장학", "기타(끝난 사역)", "미정"
F_PRAISE, F_SCHOOL, F_MOK, F_ORG = "찬양", "교회학교", "목양", "기관"
# 부설기관 계열 안(중분류) — 친구 답(2026-10-06): 복지재단과 부설기관을 「부설기관」 하나로
M_WELF, M_LIB, M_SENIOR, M_LOVE = "복지재단·봉사센터", "도서관", "시니어학교", "사랑부·샬롬부"

FAM_BY_DEPT = {
    "전도부": F_EVAN, "기도사역부": F_EVAN,
    "교회봉사부": F_SERVE, "봉사부": F_SERVE,
    "희망의복지재단": F_AFFIL, "희망의복지재단/사회봉사센터": F_AFFIL, "복지재단사업": F_AFFIL, "사회봉사센터사업": F_AFFIL, "부설기관": F_AFFIL,
    "늘푸른학교": F_AFFIL,
    "예배부": F_WORSHIP,
    "시설관리부": F_FACIL, "시설부": F_FACIL,
    "총무부": F_ADMIN, "홍보부": F_ADMIN,
    "새가족부": F_NEW, "새가족양육부": F_NEW, "양육부": F_DISC,
    "가정행복부": F_DISC, "행복사역부": F_DISC, "가정사역": F_DISC, "제자양육부": F_DISC,
    "선교부": F_MISSION, "차량부": F_CAR,
    "멀티미디어운영부": F_MEDIA, "방송운영부": F_MEDIA, "방송전산부": F_MEDIA, "새물결부": F_MEDIA,
    "문화사역부": F_CULT, "스포츠선교부": F_CULT, "스포츠운영부": F_CULT, "문화스포츠부": F_CULT,
    "재정부": F_FIN, "새성전건축위원회": F_FIN,
    "감사위원회": F_AUDIT,
    "CYLT": F_ENDED,
    "": F_UNKNOWN,
}
SCHOOL_DEPTS = ("교회학교", "교육위원회")
PRAISE_DEPTS = ("찬양부", "찬양대")

# 표준 팀 — 계열마다 (정규식, 표준 이름). 위에서부터 처음 맞는 것. 안 맞으면 바로잡은 팀 이름 그대로.
SYN = {
    F_SERVE: [(r"식당|오병이어", "식당(오병이어)"), (r"^상례", "상례"), (r"카페", "카페운영"), (r"결혼예식", "결혼예식"), (r"^데코", "데코"),
              (r"^조경", "조경"), (r"행정운영", "운영")],
    F_AFFIL: [(r"도서", "도서관(책마을)"), (r"푸드뱅크|반찬", "희망푸드뱅크"), (r"사랑의봉사|사랑의쌀|소년소녀", "사랑의봉사"), (r"그루터기", "그루터기"),
              (r"봉사센터운영", "봉사센터운영"), (r"재단운영|재단사무국|행정지원|^홍보$", "재단운영"), (r"보호|데이케어", "주야간보호"),
              (r"방과후", "방과후교실"), (r"자원봉사지원", "자원봉사지원"), (r"재가", "재가복지")],
    F_WORSHIP: [(r"본당영접", "본당영접"), (r"출입문영접", "출입문영접"), (r"웰컴|^영접|^안내$", "웰컴(영접)"), (r"세례|성찬|성례", "성례(세례·성찬)"),
                (r"기도", "기도회운영"), (r"이단", "이단대응"), (r"^(운영|기획|예배)$", "운영")],
    F_EVAN: [(r"전도대", "전도대"), (r"전도학교", "전도학교"), (r"중보기도학교", "중보기도학교"), (r"중보", "중보기도"), (r"기도실", "기도실운영"),
             (r"^운영", "운영")],
    F_MISSION: [(r"세계선교", "세계선교"), (r"국내선교", "국내선교"), (r"다문화", "다문화선교")],
    F_CAR: [(r"주차", "주차안내"), (r"차량운행", "차량운행")],
    F_MEDIA: [(r"새물결", "새물결"), (r"^방송", "방송실운영"), (r"IT|전산", "전산"), (r"홈페이지", "홈페이지운영"), (r"영상", "영상제작"),
              (r"미디어", "미디어홍보")],
    F_CULT: [(r"동호", "동호회"), (r"문화선교", "문화선교"), (r"문화행사", "문화행사"), (r"평생교육원|꿈샘", "문화교실(평생교육원)"),
             (r"스포츠", "스포츠"), (r"^운영", "운영")],
    F_FIN: [(r"계수", "계수"), (r"회계|재무", "회계·재무")],
    F_FACIL: [(r"안전", "안전관리"), (r"시설|영선|자재", "시설관리"), (r"^데코", "데코"), (r"회계", "운영")],
    F_ADMIN: [(r"대외협력", "대외협력"), (r"홍보|문서|우물가|신문|영상|^부장$", "홍보·역사"), (r"경축", "경축")],
    F_NEW: [(r"영접", "새가족영접"), (r"정착|새가족양육", "새가족정착(양육)"), (r"이단", "이단대응"), (r"^운영", "운영")],
    F_DISC: [(r"마더와이즈", "마더와이즈"), (r"하프타임", "5060하프타임"), (r"아버지학교", "아버지학교"), (r"어머니학교", "어머니학교"),
             (r"젊은부부", "젊은부부"), (r"가정사역", "가정사역"), (r"육아", "육아지원"), (r"영대디", "영대디캠프"), (r"신앙운동", "신앙운동"),
             (r"듀나미스", "듀나미스성경사관학교"), (r"학사|사역배치|^지원$", "학사·지원"), (r"양육", "제자양육"), (r"^운영", "운영")],
}


def syn(fam, tk):
    for rx, std in SYN.get(fam, []):
        if re.search(rx, tk):
            return std
    return tk or "(팀 이름 없음)"


def std_praise(tk):
    m = re.match(r"^(시온|가브리엘|임마누엘|할렐루야|글로리아|새하늘|실로암)(찬양대)?$", tk)
    if m:
        return m.group(1) + "찬양대", "찬양대"
    if tk == "찬양대중1개":
        return "찬양대(어느 찬양대인지 안 적힘)", "찬양대"
    T = "예배 찬양팀"
    if re.search(r"주일.*2부", tk): return "주일2부찬양", T
    if re.search(r"주일.*3부", tk): return "주일3부찬양", T
    if re.search(r"주일.*오후", tk): return "주일오후찬양", T
    if re.search(r"^주일(찬양)?예배찬양$", tk): return "주일찬양(2018년에 나뉘기 전)", T
    if "주일찬양예배,수요" in tk: return "찬양팀(여러 예배)", T
    if re.search(r"수요.*오전", tk): return "수요오전찬양", T
    if re.search(r"수요.*오후", tk): return "수요오후찬양", T
    if re.search(r"^수요(부흥)?기도회(찬양)?$", tk): return "수요찬양(2018년에 나뉘기 전)", T
    if tk.startswith("금요"): return "금요찬양", T
    if "찬양팀" in tk or "찬양단" in tk: return tk, T
    if "반주" in tk: return "반주", T
    return tk or "(팀 이름 없음)", "그 밖의 팀"


def std_school(tk):
    if re.search(r"어와나|꿈너머", tk): return "어와나", "프로그램"
    if re.search(r"토요드림", tk): return "토요드림학교", "프로그램"
    if re.search(r"성품스쿨|청포도|성취포상", tk): return tk, "프로그램"
    if re.search(r"전도", tk): return "교회학교전도", "전도"
    if re.search(r"영아|유아|유치|새싹", tk): return tk, "미취학"
    if re.search(r"유년|소년|초등|학년|^아동$", tk): return tk, "아동"
    if re.search(r"중등|고등", tk): return tk, "청소년"
    if tk == "청년부": return tk, "청년부"
    return "운영·지원", "운영·지원"


def std_mok(dk, tk):
    base = "젊은부부" if dk == "젊은부부" else dk[:4]          # L-12 · M-12 · C-12
    if "리더" in tk: return base + " 리더", "리더"
    if re.search(r"목양지원|목양관리", tk): return "목양지원", "지원·운영"
    if "정착관리" in tk: return "정착관리", "지원·운영"
    if tk == "젊은부부": return "젊은부부", "지원·운영"
    return "운영·행정", "지원·운영"


# ── 1차에서 여쭌 17가지(글은 여쭌 그대로) ───────────────────────────────
ASK = collections.OrderedDict([
    ("Q01", "L-12 · M-12 · C-12 줄(팀 이름이 「리더」)을 목장 리더 조직으로 보고 「따로 · 목양」에 두었습니다. 맞나요?"),
    ("Q02", "교회학교 아래 「도서관리」(2010~2012)는 2013년 부설기관 교회도서관으로 사람이 옮겨 가서 「그 밖」의 도서관으로 두었습니다. 교회학교에 둘까요?"),
    ("Q03", "「샬롬부」(2019년까지)와 「사랑부」(2024년부터)는 같은 부서인가요? 사람은 거의 이어지지 않습니다(사이가 4년 비어 있음). 지금은 따로 두었습니다."),
    ("Q04", "경로대학(2015 · 복지재단) → 늘푸른학교(2016~2019) → 은빛시니어학교(2020 · 2023~ · 교육위원회) 그리고 교육위원회 「장년부」(2022)를 한 계열 「시니어학교」로 이었습니다. 맞나요?"),
    ("Q05", "「젊은부부」 부서의 리더 8줄(2011)을 목양에 두었습니다."),
    ("Q06", "CYLT 「요셉학교」(2010~2011)와 부설기관 「다음세대리더십센터」(2013)는 무엇이었나요? 지금은 끝난 사역으로 두었습니다."),
    ("Q07", "2022년 「주보간지」 15줄은 부서 칸이 비어 있습니다. 어느 부서인가요?"),
    ("Q08", "홍보부(2010~2013)를 총무부의 홍보·역사(2014~)와 한 계열로 이었습니다. 다만 사람은 이어지지 않습니다 — 일이 같아서 이은 것입니다."),
    ("Q09", "기도사역부(2013 · 팀장 3줄)를 전도부 중보기도와 한 계열로 이었습니다. 사람은 이어지지 않습니다."),
    ("Q10", "가정행복부(2010~2013 아버지학교·어머니학교) · 행복사역부(2018~2023) · 제자양육부(2024~)를 한 계열 「제자양육·가정사역」으로 이었습니다. 가정사역을 따로 볼까요?"),
    ("Q11", "교육위원회 아래 적힌 「금요성령집회찬양」(2021 · 1줄)·「베데스다」(2021~2022 · 2줄)는 찬양으로 보았습니다(찬양부에 같은 팀이 있음)."),
    ("Q12", "새성전건축위원회 「건축헌금계수」(2012 · 4줄)를 재정에 두었습니다."),
    ("Q13", "찬양부 아래 「중보기도팀」(2024~2026)은 부서가 찬양부라 찬양에 그대로 두었습니다. 전도·중보기도로 옮길까요?"),
    ("Q14", "교회봉사부 「식당운영」(2024년까지)과 「오병이어 1·2팀」(2025년부터)을 같은 일로 이었습니다."),
    ("Q15", "문화사역부 「평생교육원」과 문화스포츠부 「꿈샘문화교실」(2022~)을 같은 일로 이었습니다. 사람 근거는 약합니다."),
    ("Q16", "양육부·새가족양육부 안의 「양육 · 학사 · 신앙운동」 팀은 부서가 새가족양육부였어도 「제자양육·가정사역」 계열로 이었습니다(2024년에 제자양육부로 감)."),
    ("Q17", "2019년 전도부 「세계선교」 1줄은 선교로 보았습니다(부서를 잘못 적은 듯)."),
])
# 친구 답(2026-10-06) — (엑셀에 적은 답 · 뜻이 둘로 읽혀 되물은 것은 「→」 뒤에, 이렇게 반영했다)
ANSWER = {
    "Q01": ("네", "그대로 — 「따로 · 목양」."),
    "Q02": ("부설기관 → 되물음: 부설기관과 복지재단을 부설기관으로 통일",
            "도서관(어린이도서관 · 교회도서관 · 책마을도서관)을 17년 내내 「부설기관」 계열에 두었습니다. 복지재단·봉사센터 계열도 「부설기관」으로 합쳤습니다."),
    "Q03": ("부설기관 → 되물음: 둘 다 부설기관",
            "샬롬부와 사랑부를 교회학교에서 빼서 「부설기관」에 두었습니다. 같은 부서인지는 아래 남은 확인 N01."),
    "Q04": ("부설기관", "경로대학 · 늘푸른학교 · 은빛시니어학교 · 장년부를 「부설기관」 계열의 시니어학교로 두었습니다."),
    "Q05": ("네", "그대로 — 목양."),
    "Q06": ("네", "그대로 — 끝난 사역."),
    "Q07": ("예배", "주보간지 15줄을 「예배」 계열로 옮겼습니다."),
    "Q08": ("네", "그대로 — 총무·홍보."),
    "Q09": ("전도대", "계열은 그대로 「전도·중보기도」입니다. 팀 이름(중보기도 · 중보기도학교 · 운영)은 명단에 적힌 대로 두었습니다."),
    "Q10": ("네 → 되물음: 한 계열로 둔다", "그대로 — 「제자양육·가정사역」 한 계열."),
    "Q11": ("네", "그대로 — 찬양."),
    "Q12": ("네", "그대로 — 재정."),
    "Q13": ("네 → 되물음: 전도·중보기도로 옮긴다",
            "찬양부 아래 중보기도팀을 「전도·중보기도」 계열로 옮겼습니다. 전도부 중보기도와는 다른 팀 「중보기도팀(찬양부)」로 표시합니다."),
    "Q14": ("네", "그대로 — 식당(오병이어)."),
    "Q15": ("네", "그대로 — 문화교실(평생교육원)."),
    "Q16": ("네", "그대로 — 제자양육·가정사역."),
    "Q17": ("네", "그대로 — 선교."),
}
# 남은 확인 — 2차에서 새로 여쭙는 것
OPEN = collections.OrderedDict([
    ("N01", "샬롬부(2019년까지)와 사랑부(2024년부터)를 둘 다 부설기관에 두었습니다. 같은 부서라면 한 팀으로 이어서 그립니다. 같은 부서인가요?"),
    ("N02", "명단이 온전해 보이지 않는 해가 있습니다(「큰 분류와 해」 시트의 노란 칸 — 2021년 전체 · 교회학교 2014·2023·2024년 · 2022년부터의 재정부). 그 해 명단에 무엇을 실었는지 사역 담당께 확인이 필요합니다."),
    ("N03", "2026년 명단에 부서와 팀이 모두 빈 줄이 하나 있습니다. 어느 부서인가요?"),
])
OK = "(친구 확인 %s)" % ANSWERED_ON


def classify(dept, team):
    """(큰 분류, 계열, 중분류, 표준 팀, 메모, 1차 물음 번호, 남은 확인 번호)"""
    dk, tk = dept_key(dept), team_key(team)

    def R(g, fam, mid, std, memo="", ask="", open_=""):
        return g, fam, mid, std, memo, ask, open_
    # 따로 — 기관 · 목양
    if re.search(r"(여전도회|남선교회|권사회|안수집사회|남선교회연합회)$", dk):
        return R(G_ORG, F_ORG, re.sub(r"^\d+", "", dk), "회장", "2010~2011년 명단에만 있다")
    if re.match(r"^[LMC]-12", dk):
        std, mid = std_mok(dk, tk)
        return R(G_MOK, F_MOK, mid, std, "", "Q01")
    if dk == "젊은부부":
        std, mid = std_mok(dk, tk)
        return R(G_MOK, F_MOK, mid, std, "", "Q05")
    # 팀 이름이 정하는 것(부서보다 먼저)
    if "찬양대" in tk and dk not in PRAISE_DEPTS:
        std, mid = std_praise(tk)
        return R(G_PRAISE, F_PRAISE, mid, std, "부서는 %s — 팀 이름이 찬양대라 찬양으로" % dk)
    if dk in SCHOOL_DEPTS:
        if tk in ("금요성령집회찬양", "베데스다"):
            std, mid = std_praise(tk)
            return R(G_PRAISE, F_PRAISE, mid, std, "부서는 교육위원회로 적혀 있다 — 찬양으로 " + OK, "Q11")
        if re.search(r"은빛시니어|장년부", tk):
            return R(G_ETC, F_AFFIL, M_SENIOR, "시니어학교", "부서는 교육위원회 — 부설기관으로 " + OK, "Q04")
        if "장학" in tk:
            return R(G_ETC, F_SCHOLAR, "", "장학", "부서는 %s — 교회학교 교사가 아니라 그 밖으로(친구 결정)" % dk)
        if "도서" in tk:
            return R(G_ETC, F_AFFIL, M_LIB, "도서관(책마을)", "부서는 %s — 부설기관으로 %s · 2013년 부설기관 교회도서관으로 사람이 이어진다" % (dk, OK), "Q02")
        if re.search(r"샬롬|사랑부", tk):
            return R(G_ETC, F_AFFIL, M_LOVE, "샬롬부" if "샬롬" in tk else "사랑부", "부서는 %s — 부설기관으로 %s" % (dk, OK), "Q03", "N01")
        std, mid = std_school(tk)
        return R(G_SCHOOL, F_SCHOOL, mid, std)
    if dk in PRAISE_DEPTS:
        if tk == "중보기도팀":
            return R(G_ETC, F_EVAN, "", "중보기도팀(찬양부)", "부서는 찬양부 — 전도·중보기도로 " + OK, "Q13")
        std, mid = std_praise(tk)
        return R(G_PRAISE, F_PRAISE, mid, std)
    # 그 밖 — 팀이 부서와 다른 계열로 가는 것
    if "경로대학" in tk:
        return R(G_ETC, F_AFFIL, M_SENIOR, "시니어학교", "2016년 늘푸른학교로 이어진다", "Q04")
    if dk == "늘푸른학교":
        return R(G_ETC, F_AFFIL, M_SENIOR, "시니어학교", "", "Q04")
    if dk == "총무부" and re.search(r"시설|영선|안전", tk):
        return R(G_ETC, F_FACIL, "", syn(F_FACIL, tk), "부서는 총무부 — 시설·안전 일은 2024년 시설부로 나뉨")
    if dk == "전도부" and "새가족" in tk:
        return R(G_ETC, F_NEW, "", syn(F_NEW, tk), "부서는 전도부(2010) — 2011년 양육부 새가족영접으로 옮겨 감")
    if dk == "전도부" and "세계선교" in tk:
        return R(G_ETC, F_MISSION, "", syn(F_MISSION, tk), "부서는 전도부로 적혀 있다 — 선교로 " + OK, "Q17")
    if dk == "양육부" and "새가족" in tk:
        return R(G_ETC, F_NEW, "", syn(F_NEW, tk), "부서는 양육부 — 2013년 새가족부로 나뉨")
    if dk == "새가족양육부" and re.search(r"학사|신앙운동|듀나미스|가정사역", tk):
        return R(G_ETC, F_DISC, "", syn(F_DISC, tk), "부서는 새가족양육부 — 양육·학사 일은 2024년 제자양육부로 감", "Q16")
    if dk == "부설기관" and "다음세대" in tk:
        return R(G_ETC, F_ENDED, "", "다음세대리더십센터", "부서는 부설기관(2013) — 끝난 사역으로 " + OK, "Q06")
    if dk == "부설기관" and "카페" in tk:
        return R(G_ETC, F_SERVE, "", syn(F_SERVE, tk), "부서는 부설기관(2013) — 2014년부터 교회봉사부 카페운영")
    if dk == "" and tk == "주보간지":
        return R(G_ETC, F_WORSHIP, "", "주보간지", "부서 칸이 비어 있다 — 예배로 " + OK, "Q07")
    fam = FAM_BY_DEPT.get(dk)
    if fam is None:
        return R(G_ETC, F_UNKNOWN, "", tk or "(팀 이름 없음)", "규칙에 없는 부서 — 계열을 정해 주세요", "", "N03")
    if fam == F_UNKNOWN:
        return R(G_ETC, F_UNKNOWN, "", tk or "(팀 이름 없음)", "부서와 팀이 모두 비어 있다", "", "N03")
    std = "새물결" if dk == "새물결부" else "스포츠" if dk in ("스포츠선교부", "스포츠운영부") else "감사" if fam == F_AUDIT else syn(fam, tk)
    mid = (M_LIB if std == "도서관(책마을)" else M_WELF) if fam == F_AFFIL else ""
    ask = {"CYLT": "Q06", "홍보부": "Q08", "기도사역부": "Q09", "새성전건축위원회": "Q12"}.get(dk, "")
    if dk in ("가정행복부", "행복사역부", "가정사역"): ask = "Q10"
    if dk == "양육부": ask = "Q16"
    if fam == F_SERVE and std == "식당(오병이어)": ask = "Q14"
    if fam == F_CULT and std == "문화교실(평생교육원)": ask = "Q15"
    if fam == F_AFFIL and std == "도서관(책마을)": ask = "Q02"
    return R(G_ETC, fam, mid, std, "", ask)


# ── 통폐합·옮김 — 사람 흐름으로 확인할 것들 ──────────────────────────────
# (해, 무슨 일, 종류, 앞 (해, 부서 정규식, 팀 정규식), 뒤 (…), 확인 번호 — Q.. 는 친구가 확인한 것 · N.. 은 남은 확인 · "" 은 여쭙지 않은 것)
EVENTS = [
    (2011, "새가족영접 팀: 전도부 → 양육부", "팀이 옮김", (2010, r"^전도부$", r"새가족"), (2011, r"^양육부$", r"새가족영접"), ""),
    (2012, "교회학교 → 교육위원회", "이름 바뀜", (2011, r"^교회학교$", r"."), (2012, r"^교육위원회$", r"."), ""),
    (2012, "새하늘찬양대: 찬양부 → 교회봉사부", "팀이 옮김", (2011, r"^찬양부$", r"새하늘"), (2012, r"^교회봉사부$", r"새하늘"), ""),
    (2012, "희망의복지재단 → 복지재단사업 + 사회봉사센터사업(이 해만 나뉘어 적힘)", "나뉨", (2011, r"^희망의복지재단$", r"."), (2012, r"^(복지재단사업|사회봉사센터사업|희망의복지재단)$", r"."), ""),
    (2013, "양육부에서 새가족부가 나뉨", "나뉨", (2012, r"^양육부$", r"새가족"), (2013, r"^새가족부$", r"."), ""),
    (2013, "도서관: 교육위원회 도서관리 + 복지재단사업 어린이도서관 → 부설기관", "합침 · 팀이 옮김", (2012, r".", r"도서"), (2013, r"^부설기관$", r"도서"), "Q02"),
    (2013, "멀티미디어운영부 → 방송운영부", "이름 바뀜", (2012, r"^멀티미디어운영부$", r"."), (2013, r"^방송운영부$", r"."), ""),
    (2013, "시설관리부가 새로 생김(총무부의 시설·영선과 그 해 나란히 있음)", "새로 생김", (2012, r"^총무부$", r"시설|영선"), (2013, r"^시설관리부$", r"."), ""),
    (2014, "새가족부 + 양육부 → 새가족,양육부", "합침", (2013, r"^(새가족부|양육부)$", r"."), (2014, r"^새가족양육부$", r"."), ""),
    (2014, "부설기관 → 희망의복지재단(책마을도서관)", "합침", (2013, r"^부설기관$", r"."), (2014, r"^희망의복지재단$", r"."), "Q02"),
    (2014, "방송운영부 → 방송전산부", "이름 바뀜", (2013, r"^방송운영부$", r"."), (2014, r"^방송전산부$", r"."), ""),
    (2014, "홍보부 → 총무부 홍보", "합침", (2013, r"^홍보부$", r"."), (2015, r"^총무부$", r"홍보"), "Q08"),
    (2014, "기도사역부 → 전도부 중보기도", "합침", (2013, r"^기도사역부$", r"."), (2014, r"^전도부$", r"중보"), "Q09"),
    (2014, "L-12운영부 → L-12(목양지원·운영)", "합침", (2013, r"^L-12운영부$", r"."), (2014, r"^L-12$", r"목양지원|운영"), ""),
    (2014, "글로리아찬양대가 없어짐 — 대원이 다른 찬양대로", "없어짐", (2013, r"^찬양", r"글로리아"), (2014, r"^찬양부$", r"찬양대"), ""),
    (2014, "할렐루야찬양대가 없어짐 — 2024년에 다시 생김", "없어짐", (2013, r"^찬양", r"할렐루야"), (2014, r"^찬양부$", r"찬양대"), ""),
    (2015, "교회봉사부 → 봉사부", "이름 바뀜", (2014, r"^교회봉사부$", r"."), (2015, r"^봉사부$", r"."), ""),
    (2016, "경로대학(복지재단) → 늘푸른학교(부서로)", "팀이 옮김", (2015, r"복지", r"경로대학"), (2016, r"^늘푸른학교$", r"."), "Q04"),
    (2016, "꿈너머학교 → 어와나", "이름 바뀜", (2015, r"^교육위원회$", r"꿈너머"), (2016, r"^교육위원회$", r"꿈너머|어와나"), ""),
    (2017, "스포츠선교부 → 스포츠운영부", "이름 바뀜", (2016, r"^스포츠선교부$", r"."), (2017, r"^스포츠운영부$", r"."), ""),
    (2018, "봉사부 → 교회봉사부", "이름 바뀜", (2017, r"^봉사부$", r"."), (2018, r"^교회봉사부$", r"."), ""),
    (2018, "시설관리부 → 총무부 시설관리", "합침", (2017, r"^시설관리부$", r"."), (2018, r"^총무부$", r"시설"), ""),
    (2018, "행복사역부가 생김 — 가정행복부(2013년까지)와 사람은 이어지지 않음", "새로 생김", (2013, r"^가정행복부$", r"."), (2018, r"^행복사역부$", r"."), "Q10"),
    (2020, "문화사역부 + 스포츠운영부 → 문화스포츠부", "합침", (2019, r"^(문화사역부|스포츠운영부)$", r"."), (2020, r"^문화스포츠부$", r"."), ""),
    (2020, "새물결부 → 방송전산부 새물결", "합침", (2019, r"^새물결부$", r"."), (2020, r"^방송전산부$", r"새물결"), ""),
    (2020, "늘푸른학교 → 교육위원회 은빛시니어학교", "팀이 옮김", (2019, r"^늘푸른학교$", r"."), (2020, r"^교육위원회$", r"은빛"), "Q04"),
    (2023, "교육위원회 장년부(2022) → 은빛시니어학교", "이름 바뀜", (2022, r"^교육위원회$", r"장년부"), (2023, r"^교육위원회$", r"은빛"), "Q04"),
    (2024, "새가족양육부 → 새가족부", "이름 바뀜", (2023, r"^새가족양육부$", r"영접|정착|운영"), (2024, r"^새가족부$", r"."), ""),
    (2024, "행복사역부 → 제자양육부", "합침", (2023, r"^(행복사역부|가정사역)$", r"."), (2024, r"^제자양육부$", r"."), "Q10"),
    (2024, "새가족양육부의 학사·신앙운동 → 제자양육부", "나뉨", (2023, r"^새가족양육부$", r"학사|신앙"), (2024, r"^제자양육부$", r"."), "Q16"),
    (2024, "총무부 시설관리·안전관리 → 시설부", "나뉨", (2023, r"^총무부$", r"시설|안전"), (2024, r"^시설부$", r"."), ""),
    (2024, "이단대응 팀: 예배부 → 새가족부", "팀이 옮김", (2023, r"^예배부$", r"이단"), (2024, r"^새가족부$", r"이단"), ""),
    (2024, "전도대 → 행복전도대(금·토)", "이름 바뀜", (2023, r"^전도부$", r"전도대"), (2024, r"^전도부$", r"행복전도대"), ""),
    (2024, "샬롬부(2019년까지) … 사랑부(2024년부터)", "이어지는지 모름", (2019, r"^교육위원회$", r"샬롬"), (2025, r"^교육위원회$", r"사랑부"), "N01"),
    (2025, "식당운영 → 오병이어 1·2팀", "이름 바뀜", (2024, r"봉사부$", r"식당"), (2025, r"봉사부$", r"오병이어"), "Q14"),
    (2026, "희망의복지재단 → 희망의복지재단/사회 봉사센터", "이름 바뀜", (2025, r"^희망의복지재단$", r"."), (2026, r"^희망의복지재단/사회봉사센터$", r"."), ""),
]


def span(years):
    ys = sorted(set(years)); out = []; a = b = ys[0]
    for y in ys[1:]:
        if y == b + 1: b = y
        else: out.append((a, b)); a = b = y
    out.append((a, b))
    return " · ".join(str(a) if a == b else "%d~%d" % (a, b) for a, b in out)


def load(paths):
    rows = []
    for p in paths:
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        it = wb.worksheets[0].iter_rows(values_only=True)
        head = [nfc(h) for h in next(it)]
        ix = {h: i for i, h in enumerate(head)}
        for need in ("년도", "부서", "팀명", "이름", "목장"):
            if need not in ix: raise SystemExit("칸이 없다: %s (%s)" % (need, p))
        for r in it:
            if not any(r): continue
            y = re.match(r"(\d{4})", nfc(r[ix["년도"]]))
            if not y: continue
            rows.append({"year": int(y.group(1)), "dept": nfc(r[ix["부서"]]), "team": nfc(r[ix["팀명"]]),
                         "name": ns(nfc(r[ix["이름"]])), "mok": ns(nfc(r[ix["목장"]]))})
    return rows


# ── 엑셀 꾸미기 ─────────────────────────────────────────────────────
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
HEAD_FONT = Font(bold=True, color="FFFFFF")
ASK_FILL = PatternFill("solid", fgColor="FFF2CC")
GROUP_FILL = {G_PRAISE: "E2EFDA", G_SCHOOL: "DDEBF7", G_ETC: "FFFFFF", G_MOK: "EDEDED", G_ORG: "EDEDED"}
WRAP = Alignment(wrap_text=True, vertical="top")


def sheet(wb, title, head, widths, freeze="A2"):
    ws = wb.create_sheet(title)
    ws.append(head)
    for c in ws[1]:
        c.fill, c.font, c.alignment = HEAD_FILL, HEAD_FONT, Alignment(wrap_text=True, vertical="center")
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = freeze
    ws.page_setup.orientation = "landscape"          # 찍어 볼 때 — 가로 · 한 장 너비에 맞춤
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr = openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.print_title_rows = "1:1"
    return ws


def finish(ws, wrap_cols=(), filt=True):
    if filt and ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    top = Alignment(vertical="top")            # 줄이 여러 줄로 접혀도 같은 줄의 칸들이 위로 맞게
    for row in ws.iter_rows(min_row=2):
        for i, c in enumerate(row, 1):
            c.alignment = WRAP if i in wrap_cols else top


def paint(row, fill=ASK_FILL):
    for c in row:
        c.fill = fill


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")       # 윈도 콘솔(cp949)에서 「—」 같은 글자로 멈추지 않게
    args = sys.argv[1:]
    out = DEFAULT_OUT
    if "--out" in args:
        i = args.index("--out"); out = args[i + 1]; del args[i:i + 2]
    paths = args or DEFAULT_IN
    rows = load(paths)
    years = sorted({r["year"] for r in rows})
    for r in rows:
        r["g"], r["fam"], r["mid"], r["std"], r["memo"], r["ask"], r["open"] = classify(r["dept"], r["team"])
        r["dk"], r["tk"] = dept_key(r["dept"]), team_key(r["team"])

    # ── 쌍(원래 부서, 원래 팀) ──
    pair = collections.OrderedDict()
    for r in rows:
        k = (r["dept"], r["team"])
        p = pair.setdefault(k, {"years": collections.Counter(), **{x: r[x] for x in ("g", "fam", "mid", "std", "memo", "ask", "open", "dk")}})
        p["years"][r["year"]] += 1

    def people(sel):
        y, drx, trx = sel
        return {r["name"] for r in rows if r["year"] == y and re.search(drx, r["dk"]) and re.search(trx, r["tk"])}

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # ① 읽는 법
    ws = sheet(wb, "읽는 법", ["부서 이음표 2차 — 읽는 법"], [118], freeze=None)
    ws.page_setup.fitToHeight = 1          # 찍으면 한 장에
    guide = [
        "이 엑셀은 2010~2026년 사역 임명 명단 %s줄에 적힌 부서·팀 이름을 한 틀로 잇는 표의 2차입니다." % format(len(rows), ","),
        "1차에서 여쭌 %d가지에 답해 주신 것(%s)을 반영했습니다. 남은 확인은 %d가지입니다." % (len(ASK), ANSWERED_ON, len(OPEN)),
        "",
        "■ 세 겹으로 잇습니다",
        "   큰 분류 — 찬양 · 교회학교 · 그 밖. 목장 리더 줄(L-12 등)과 기관 줄(여전도회 등)은 「따로」 둡니다.",
        "   계열 — 이름이 바뀌거나 합치고 나뉜 부서를 한 선으로 이어 보는 단위입니다. 예) 방송·전산 = 멀티미디어운영부 → 방송운영부 → 방송전산부(+새물결부).",
        "   표준 팀 — 띄어쓰기·오타·이름 바뀜을 모은 팀 이름입니다. 예) 식당운영 · 오병이어 1팀 → 식당(오병이어).",
        "",
        "■ 1차와 달라진 것",
        "   「부설기관」 계열을 새로 두고, 복지재단·봉사센터를 거기로 합쳤습니다. 도서관 · 시니어학교(경로대학·늘푸른학교·은빛시니어학교) · 샬롬부 · 사랑부가 여기에 있습니다.",
        "   샬롬부와 사랑부는 교회학교에서 빠졌습니다. 찬양부 아래 중보기도팀은 전도·중보기도로, 부서 칸이 빈 「주보간지」는 예배로 옮겼습니다.",
        "",
        "■ 시트(왼쪽부터)",
        "   확인 — 1차의 %d가지와 주신 답, 그렇게 반영한 내용. 맨 아래 노란 줄이 남은 확인 %d가지입니다." % (len(ASK), len(OPEN)),
        "   큰 분류와 해 — 해마다 찬양·교회학교·그 밖 자리 수와 사람 수, 명단이 온전해 보이지 않는 해.",
        "   분류 안 — 찬양(찬양대 · 예배 찬양팀), 교회학교(미취학 · 아동 · 청소년), 부설기관(복지재단·봉사센터 · 도서관 · 시니어학교 · 사랑부·샬롬부)을 한 번 더 나눈 해별 자리 수.",
        "   계열 한눈에 — 계열마다 해별 자리 수. 선이 끊기거나 갑자기 튀는 해가 있으면 이음이 틀렸을 수 있습니다.",
        "   통폐합·옮김 — 부서가 합치고 나뉘고, 팀이 부서를 옮긴 일. 앞뒤 해에 같은 분이 몇 분 이어졌는지를 근거로 적었습니다.",
        "   부서·팀 이음표 — 원래 부서·팀 이름 %d가지가 어디로 이어지는지. 고칠 때는 이 시트의 맨 오른쪽 칸에 적어 주세요." % len(pair),
        "   팀 한눈에 — 표준 팀마다 해별 자리 수.",
        "",
        "■ 읽을 때",
        "   자리 = 명단 한 줄(한 분이 그 해 한 팀에 임명된 것). 한 분이 여러 자리를 맡으면 자리는 여러 개입니다.",
        "   사람 수는 어림입니다 — 이름과 목장이 같으면 한 분으로 셌습니다. 실제 통계는 교인 번호로 다시 셉니다.",
        "   노란 줄 = 아직 확인이 남은 줄.",
        "   이 엑셀에는 사람 이름이 없습니다.",
    ]
    for g in guide:
        ws.append([g])
    for row in ws.iter_rows(min_row=2):
        row[0].alignment = WRAP

    # ② 확인 — 1차 물음과 답 · 남은 확인
    ws = sheet(wb, "확인", ["번호", "여쭌 것", "주신 답(%s)" % ANSWERED_ON, "이렇게 반영했습니다 / 남은 확인은 답을 적어 주세요", "걸린 줄(자리)"],
               [8, 78, 34, 62, 12])
    for q, text in ASK.items():
        ans, did = ANSWER[q]
        ws.append([q, text, ans, did, sum(1 for r in rows if r["ask"] == q)])
    ws.append(["", "■ 남은 확인", "", "", None])
    ws[ws.max_row][1].font = Font(bold=True)
    for q, text in OPEN.items():
        n = sum(1 for r in rows if r["open"] == q)
        ws.append([q, text, "", "", n or None])
        paint(ws[ws.max_row])
    finish(ws, wrap_cols=(2, 3, 4), filt=False)

    # ③ 계열 한눈에
    head = ["큰 분류", "계열"] + years + ["합", "이 계열에 든 원래 부서(해)"]
    ws = sheet(wb, "계열 한눈에", head, [13, 20] + [7] * len(years) + [9, 90], freeze="C2")
    fam_year = collections.defaultdict(collections.Counter)
    fam_depts = collections.defaultdict(lambda: collections.defaultdict(set))
    for r in rows:
        fam_year[(r["g"], r["fam"])][r["year"]] += 1
        fam_depts[(r["g"], r["fam"])][r["dept"] or "(부서 빈칸)"].add(r["year"])
    order = sorted(fam_year, key=lambda k: (GROUPS.index(k[0]), -sum(fam_year[k].values())))
    for k in order:
        depts = sorted(fam_depts[k].items(), key=lambda kv: min(kv[1]))
        if k[1] == F_ORG:
            dtext = "여전도회 · 남선교회 · 권사회 · 안수집사회 %d곳(2010~2011)" % len(depts)
        else:
            dtext = " / ".join("%s(%s)" % (d, span(ys)) for d, ys in depts)
        ws.append([k[0], k[1]] + [fam_year[k].get(y) or None for y in years] + [sum(fam_year[k].values()), dtext])
        paint(ws[ws.max_row][:2], PatternFill("solid", fgColor=GROUP_FILL[k[0]]))
    ws.append(["합", ""] + [sum(fam_year[k].get(y, 0) for k in order) for y in years] + [len(rows), ""])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    finish(ws, wrap_cols=(len(head),), filt=False)

    # ④ 통폐합·옮김
    ws = sheet(wb, "통폐합·옮김", ["해", "무슨 일", "종류", "앞(해 · 사람 수)", "뒤(해 · 사람 수)", "같은 분", "근거", "확인"],
               [7, 60, 16, 16, 16, 9, 36, 30])
    ev_out = []          # (해, 무슨 일, 종류, 앞 사람 수, 뒤 사람 수, 같은 분, 근거, 앞 글, 뒤 글, 확인 글, 노란 줄)
    for (y, what, kind, a, b, code) in EVENTS:
        A, B = people(a), people(b)
        same = len(A & B)
        base = min(len(A), len(B)) or 1
        if kind in ("없어짐",):
            why = "%d분 가운데 %d분이 다른 찬양대에" % (len(A), same)
        elif same == 0:
            why = "사람은 이어지지 않음 — 일이 같아 이음"
        elif same >= 3 and same / base >= 0.5:
            why = "뚜렷함"
        elif base < 5:
            why = "사람이 적어 판단하기 어려움"
        else:
            why = "약함"
        if code.startswith("Q"):
            note, yellow = "친구 확인 %s (%s)" % (ANSWERED_ON, code), False
        elif code.startswith("N"):
            note, yellow = "남은 확인 " + code, True
        else:
            weak = why != "뚜렷함" and kind != "없어짐"
            note, yellow = ("근거가 약함 — 틀렸으면 알려 주세요" if weak else ""), weak
        ev_out.append((y, what, kind, len(A), len(B), same, why, "%d년 %d명" % (a[0], len(A)), "%d년 %d명" % (b[0], len(B)), note, yellow))

    # 명단에서 빠진 것 — 사람이 옮긴 것이 아니라 그 해부터 명단에 싣지 않은 것으로 보이는 줄
    def seats(y, pred):
        return sum(1 for r in rows if r["year"] == y and pred(r))
    is_org = lambda r: r["g"] == G_ORG
    is_leader = lambda r: r["g"] == G_MOK and r["mid"] == "리더"
    is_fin = lambda r: r["dk"] == "재정부"
    absent = [
        (2012, "기관(여전도회·남선교회·권사회·안수집사회 회장) 줄이 명단에서 빠짐", is_org, (2011, 2012)),
        (2012, "목장 리더 줄(L-12·M-12 「리더」)이 크게 줆", is_leader, (2011, 2012)),
        (2020, "목장 리더 줄이 다시 크게 줆", is_leader, (2019, 2020)),
        (2022, "재정부 줄이 명단에서 거의 빠짐", is_fin, (2021, 2022, 2023)),
    ]
    for (y, what, pred, ys) in absent:
        text = " → ".join("%d년 %d줄" % (yy, seats(yy, pred)) for yy in ys)
        ev_out.append((y, what, "명단에서 빠짐", 0, 0, None, text, "", "", "남은 확인 N02", True))
    ev_out.sort(key=lambda e: e[0])          # 같은 해는 적은 차례 그대로(파이썬 정렬은 차례를 지킨다)
    for (y, what, kind, na, nb, same, why, ta, tb, note, yellow) in ev_out:
        ws.append([y, what, kind, ta, tb, same, why, note])
        if yellow:
            paint(ws[ws.max_row])
    finish(ws, wrap_cols=(2, 7, 8))

    # ⑤ 부서·팀 이음표
    head = ["큰 분류", "계열", "중분류", "표준 팀", "원래 부서", "원래 팀", "나온 해", "자리", "메모", "확인", "고칠 내용(적어 주세요)"]
    ws = sheet(wb, "부서·팀 이음표", head, [13, 20, 16, 26, 24, 34, 24, 7, 56, 12, 34])
    keys = sorted(pair, key=lambda k: (GROUPS.index(pair[k]["g"]), pair[k]["fam"], pair[k]["mid"], pair[k]["std"], min(pair[k]["years"]), k[0], k[1]))
    for k in keys:
        p = pair[k]
        mark = p["open"] or (p["ask"] + " 확인됨" if p["ask"] else "")
        ws.append([p["g"], p["fam"], p["mid"], p["std"], k[0] or "(빈칸)", k[1] or "(빈칸)", span(p["years"]), sum(p["years"].values()),
                   p["memo"], mark, ""])
        row = ws[ws.max_row]
        row[0].fill = PatternFill("solid", fgColor=GROUP_FILL[p["g"]])
        if p["open"]:
            paint(row[1:])
    finish(ws, wrap_cols=(9, 11))

    # ⑥ 팀 한눈에
    head = ["큰 분류", "계열", "중분류", "표준 팀"] + years + ["합"]
    ws = sheet(wb, "팀 한눈에", head, [13, 20, 16, 30] + [7] * len(years) + [9], freeze="E2")
    team_year = collections.defaultdict(collections.Counter)
    for r in rows:
        team_year[(r["g"], r["fam"], r["mid"], r["std"])][r["year"]] += 1
    for k in sorted(team_year, key=lambda k: (GROUPS.index(k[0]), k[1], k[2], -sum(team_year[k].values()))):
        ws.append(list(k) + [team_year[k].get(y) or None for y in years] + [sum(team_year[k].values())])
        ws[ws.max_row][0].fill = PatternFill("solid", fgColor=GROUP_FILL[k[0]])
    finish(ws)

    # ⑦ 분류 안 — 찬양 · 교회학교 · 부설기관을 중분류로 한 번 더 나눈 해별 자리 수
    head = ["분류", "중분류"] + years + ["합"]
    ws = sheet(wb, "분류 안", head, [18, 20] + [7] * len(years) + [9], freeze="C2")
    mid_year = collections.defaultdict(collections.Counter)
    INNER = [G_PRAISE, G_SCHOOL, "그 밖 · " + F_AFFIL]
    for r in rows:
        label = r["g"] if r["g"] in (G_PRAISE, G_SCHOOL) else ("그 밖 · " + F_AFFIL if r["fam"] == F_AFFIL else None)
        if label:
            mid_year[(label, r["mid"])][r["year"]] += 1
    for k in sorted(mid_year, key=lambda k: (INNER.index(k[0]), -sum(mid_year[k].values()))):
        ws.append(list(k) + [mid_year[k].get(y) or None for y in years] + [sum(mid_year[k].values())])
        ws[ws.max_row][0].fill = PatternFill("solid", fgColor=GROUP_FILL.get(k[0], "FCE4D6"))
    finish(ws, filt=False)

    # ⑧ 큰 분류와 해
    head = ["해", "자리 — 찬양", "교회학교", "그 밖", "따로 · 목양", "따로 · 기관", "합",
            "사람(어림) — 찬양", "교회학교", "그 밖", "세 분류 합쳐(한 분은 한 번)", "두 분류 이상에서 봉사", "명단이 온전해 보이지 않는 곳(남은 확인 N02)"]
    ws = sheet(wb, "큰 분류와 해", head, [7, 11, 10, 9, 11, 11, 9, 13, 10, 9, 13, 11, 70])
    by_group = {}
    for y in years:
        rs = [r for r in rows if r["year"] == y]
        seat = collections.Counter(r["g"] for r in rs)
        pg = collections.defaultdict(set)
        for r in rs:
            if r["g"] in (G_PRAISE, G_SCHOOL, G_ETC):
                pg[(r["name"], r["mok"])].add(r["g"])
        ppl = collections.Counter(g for s in pg.values() for g in s)
        by_group[y] = (seat, ppl, len(pg), sum(1 for s in pg.values() if len(s) >= 2), len(rs))
    sc = lambda y: by_group[y][0].get(G_SCHOOL, 0)          # 그 해 교회학교 자리
    pr = lambda y: by_group[y][0].get(G_PRAISE, 0)
    # 명단이 온전해 보이지 않는 해 — **숫자로 미루어 본 후보**다(사역 담당 확인 전). 해를 고르는 것은 사람이, 수는 여기서 센다.
    partial = {
        2014: "교회학교 %d줄 — 앞뒤 해(%d · %d)보다 적다. 부서는 다 있는데 부서마다 인원이 얇다" % (sc(2014), sc(2013), sc(2015)),
        2021: "전체 — 부서 %d곳만 실렸다. 선교부·총무부·새가족양육부·M-12가 통째로 없고 찬양이 %d줄뿐이다(2020년 %d줄)"
              % (len({r["dk"] for r in rows if r["year"] == 2021}), pr(2021), pr(2020)),
        2022: "재정부가 없다(2021년 %d줄) · 찬양 %d줄은 2020년(%d줄)의 절반"
              % (sum(1 for r in rows if r["year"] == 2021 and r["dk"] == "재정부"), pr(2022), pr(2020)),
        2023: "교회학교 %d줄 — 2022년(%d) · 2025년(%d)의 절반 안팎" % (sc(2023), sc(2022), sc(2025)),
        2024: "교회학교 %d줄 — 2022년(%d) · 2025년(%d)의 절반 이하" % (sc(2024), sc(2022), sc(2025)),
    }
    partial = {y: t for y, t in partial.items() if y in by_group}
    for y in years:
        seat, ppl, n, multi, total = by_group[y]
        ws.append([y] + [seat.get(g, 0) for g in GROUPS] + [total] + [ppl.get(g, 0) for g in (G_PRAISE, G_SCHOOL, G_ETC)]
                  + [n, multi, partial.get(y, "")])
        if y in partial:
            ws[ws.max_row][12].fill = ASK_FILL
    finish(ws, wrap_cols=(13,), filt=False)

    want = ["읽는 법", "확인", "큰 분류와 해", "분류 안", "계열 한눈에", "통폐합·옮김", "부서·팀 이음표", "팀 한눈에"]
    wb._sheets.sort(key=lambda w: want.index(w.title))          # 요약 시트를 앞으로(openpyxl 에 차례를 한 번에 정하는 길이 없다)
    os.makedirs(os.path.dirname(out), exist_ok=True)
    wb.save(out)

    # ── 화면에는 수만 ──
    print("줄", len(rows), "· 부서", len({r["dept"] for r in rows}), "· 쌍", len(pair), "· 계열", len(fam_year), "· 표준 팀", len(team_year))
    print("남은 확인 줄(쌍)", sum(1 for p in pair.values() if p["open"]), "· 미정", sum(1 for p in pair.values() if p["fam"] == F_UNKNOWN))
    print("해 | 찬양 교회학교 그밖 목양 기관 | 사람 찬양 교회학교 그밖 전체 겹침")
    for y in years:
        seat, ppl, n, multi, _total = by_group[y]
        print(y, "|", *[seat.get(g, 0) for g in GROUPS], "|", *[ppl.get(g, 0) for g in (G_PRAISE, G_SCHOOL, G_ETC)], n, multi)
    print("계열(자리 합)")
    for k in order:
        print("  %s | %s | %d | %s" % (k[0], k[1], sum(fam_year[k].values()), span(fam_year[k].keys())))
    print("분류 안(자리 합)")
    for k in sorted(mid_year, key=lambda k: (INNER.index(k[0]), -sum(mid_year[k].values()))):
        print("  %s | %s | %d | %s" % (k[0], k[1], sum(mid_year[k].values()), " ".join(str(mid_year[k].get(y, 0)) for y in years)))
    print("통폐합·옮김 근거")
    for e in ev_out:
        print("  %d %s [%s] %s→%s 같은 분 %s (%s) %s" % (e[0], e[1], e[2], e[3], e[4], e[5], e[6], e[9]))
    print("저장:", out)


if __name__ == "__main__":
    main()
