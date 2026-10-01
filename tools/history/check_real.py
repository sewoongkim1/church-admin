"""사역 이력 맞춤 규칙 대조 — 이 PC 에서만(진짜 명부·명단 · 2026-10-01).

  python tools/history/check_real.py [명단 엑셀] [교인명부 정리 엑셀]
    기본: C:/Projects/church-admin/Data/2022-2026_사역임명_통합.xlsx · C:/Projects/교인목록_2026_09_29_정리.xlsx

history-match.ts(서버와 같은 파일)를 진짜 데이터로 돌려 **수만** 찍고, 2026-10-01 독립 검증이 짚은 줄(엑셀 줄 번호)이
기대대로인지 본다. 이름·교인ID 는 찍지 않는다. 명부·명단을 옮긴 JSON 은 저장소 밖 임시 폴더에 두고 끝나면 지운다.
⚠️ 가족 근거는 이 도구만 「신앙세대주 이름 + 주소」로 묶는다(정리 엑셀에 세대주 교인ID 가 없다) — 서버는 household_id 다.
   그래서 운영 결과와 한두 줄 다를 수 있다. 기대값(4,042 / 51)은 2026-09-29 명부 기준이다.
"""
import json, os, re, shutil, subprocess, sys, tempfile, unicodedata
import openpyxl

ROSTER = sys.argv[1] if len(sys.argv) > 1 else r"C:\Projects\church-admin\Data\2022-2026_사역임명_통합.xlsx"
DIRECTORY = sys.argv[2] if len(sys.argv) > 2 else r"C:\Projects\교인목록_2026_09_29_정리.xlsx"
HERE = os.path.dirname(os.path.abspath(__file__))


def nfc(v):
    return unicodedata.normalize("NFC", "" if v is None else str(v)).strip()


def year_of(v):
    m = re.match(r"(\d{4})", nfc(v))
    return int(m.group(1)) if m and m.group(1) != "0000" else None


def month_of(v):
    m = re.match(r"\d{4}-(\d{2})", nfc(v))
    return int(m.group(1)) if m else None


def people():
    ws = openpyxl.load_workbook(DIRECTORY, read_only=True, data_only=True)["교인목록"]
    it = ws.iter_rows(values_only=True)
    ix = {h: i for i, h in enumerate(next(it))}
    out = []
    for r in it:
        if r[ix["교인ID"]] is None:
            continue
        head = nfc(r[ix["신앙세대주"]])
        out.append(dict(person_id=int(r[ix["교인ID"]]), name=nfc(r[ix["이름"]]), gender=nfc(r[ix["성별"]]), kind2=nfc(r[ix["교인구분2"]]),
                        mok1=nfc(r[ix["목장1"]]), mok3=nfc(r[ix["목장3"]]), school_dept=nfc(r[ix["교회학교부서"]]),
                        position=nfc(r[ix["직분"]]), position_detail=nfc(r[ix["직분상세"]]),
                        birth_year=year_of(r[ix["생년월일"]]), birth_month=month_of(r[ix["생년월일"]]), reg_year=year_of(r[ix["등록일"]]),
                        household=(head + "|" + nfc(r[ix["주소"]])) if head else ""))
    return out


def rows():
    ws = openpyxl.load_workbook(ROSTER, read_only=True, data_only=True).worksheets[0]
    it = ws.iter_rows(values_only=True)
    ix = {nfc(h): i for i, h in enumerate(next(it))}
    get = lambda r, k: nfc(r[ix[k]]) if k in ix and ix[k] < len(r) else ""
    out = []
    for n, r in enumerate(it, start=2):                     # 엑셀 줄 번호(머리 = 1)
        out.append(dict(id=n, year=int(get(r, "년도")), committee=get(r, "부서"), team=get(r, "팀명"), name=get(r, "이름"),
                        position=get(r, "직분"), mok=get(r, "목장"), renewal=get(r, "신규 / 유지"), link_how="auto", person_id=None))
    return out


def main():
    tmp = tempfile.mkdtemp(prefix="history-check-")       # 저장소 밖
    try:
        for name, data in (("people.json", people()), ("rows.json", rows())):
            with open(os.path.join(tmp, name), "w", encoding="utf-8") as f:   # 닫고 나서 node 가 읽는다(지울 때 잠김 없게)
                json.dump(data, f, ensure_ascii=False)
        r = subprocess.run(["node", "--experimental-strip-types", "--no-warnings", os.path.join(HERE, "check_real.mjs"), tmp])
        sys.exit(r.returncode)
    finally:
        # ⚠️ 말없이 넘어가지 않는다 — 진짜 이름·교인ID·세대주 이름과 주소가 든 JSON 이 남는다(백신·색인기가 잠깐 잡고 있으면 못 지운다).
        #    경로만 찍는다(내용은 찍지 않는다).
        try:
            shutil.rmtree(tmp)
        except OSError:
            print(f"임시 폴더를 못 지웠어요: {tmp} — 손으로 지워 주세요", file=sys.stderr)
            sys.exit(1)


if __name__ == "__main__":
    main()
